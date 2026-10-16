#!/usr/bin/env python3
"""Extract content from all uploaded PDFs and DOCX files."""
import os
import json
import pdfplumber
from docx import Document

UPLOAD_DIR = "/home/z/my-project/upload"
OUTPUT_DIR = "/home/z/my-project/extracted"
os.makedirs(OUTPUT_DIR, exist_ok=True)

files = [
    "OSW_LS3_LT_T1.pdf",
    "OSW_LS3_LT_T2.pdf",
    "OSW_LS3_LT_T3.pdf",
    "OSW_LS3_LT_T4.pdf",
    "OSW_LS3_LT_T5.pdf",
    "OSW_LS3_LT_T6.pdf",
    "OSW_LS3_LT_T6.docx",
]

def extract_pdf(path):
    out = {"pages": [], "raw_text": ""}
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            txt = page.extract_text() or ""
            # Try to extract tables
            tables = page.extract_tables() or []
            out["pages"].append({
                "page": i + 1,
                "text": txt,
                "tables": tables,
                "width": page.width,
                "height": page.height,
            })
            out["raw_text"] += f"\n\n--- Page {i+1} ---\n{txt}"
    return out

def extract_docx(path):
    doc = Document(path)
    out = {"paragraphs": [], "tables": [], "raw_text": ""}
    for p in doc.paragraphs:
        if p.text.strip():
            out["paragraphs"].append({
                "text": p.text,
                "style": p.style.name if p.style else "Normal",
            })
            out["raw_text"] += p.text + "\n"
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        out["tables"].append({"index": ti, "rows": rows})
        out["raw_text"] += "\n[TABLE " + str(ti) + "]\n" + "\n".join(
            " | ".join(r) for r in rows
        ) + "\n"
    return out

summary = {}
for f in files:
    p = os.path.join(UPLOAD_DIR, f)
    if not os.path.exists(p):
        print(f"Missing: {f}")
        continue
    print(f"Processing: {f}")
    ext = f.split(".")[-1].lower()
    try:
        if ext == "pdf":
            data = extract_pdf(p)
        elif ext == "docx":
            data = extract_docx(p)
        else:
            continue
        out_name = f.replace(".pdf", "_pdf").replace(".docx", "_docx") + ".json"
        with open(os.path.join(OUTPUT_DIR, out_name), "w") as fp:
            json.dump(data, fp, ensure_ascii=False, indent=2)
        summary[f] = {
            "pages": len(data.get("pages", [])),
            "paragraphs": len(data.get("paragraphs", [])),
            "tables": len(data.get("tables", [])),
            "text_len": len(data.get("raw_text", "")),
        }
        print(f"  -> OK ({summary[f]})")
    except Exception as e:
        print(f"  -> ERROR: {e}")
        summary[f] = {"error": str(e)}

with open(os.path.join(OUTPUT_DIR, "_summary.json"), "w") as fp:
    json.dump(summary, fp, indent=2)

print("\n=== SUMMARY ===")
print(json.dumps(summary, indent=2))
